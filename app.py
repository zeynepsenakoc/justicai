import os
from dotenv import load_dotenv

load_dotenv()

import json
import logging
from datetime import datetime
import io
import time
import random
import re
import unittest
import tests
import hashlib #
import base64  

from flask import Flask, render_template, request, session, redirect, url_for, flash, send_file, abort
from flask_mail import Mail, Message
import requests
from playwright.sync_api import sync_playwright
from werkzeug.security import generate_password_hash, check_password_hash
from flask_login import LoginManager, login_user, login_required, logout_user, current_user
from sqlalchemy import func 
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address


from docx import Document
from docx.shared import Pt
import qrcode


from config import INTERROGATOR_PROMPT, WRITER_PROMPT, CRITIC_PROMPT, FINALIZER_PROMPT, CATEGORIES, PETITION_HTML_TEMPLATE
from ocr_service import extract_text_from_file
from models import db, User, Petition, Feedback
from logic_services import search_legal_docs, check_rules, vector_db, ethics_guard, success_predictor
from privacy_service import privacy_guard

app = Flask(__name__)


app.secret_key = os.getenv("FLASK_SECRET")
if not app.secret_key:
    raise ValueError("GÜVENLİK HATASI: 'FLASK_SECRET' ortam değişkeni ayarlanmamış! .env dosyasını kontrol edin.")


app.config['MAIL_SERVER'] = os.getenv('MAIL_SERVER', 'smtp.gmail.com')
app.config['MAIL_PORT'] = int(os.getenv('MAIL_PORT', 465))  
app.config['MAIL_USE_TLS'] = os.getenv('MAIL_USE_TLS', 'False') == 'True' 
app.config['MAIL_USE_SSL'] = os.getenv('MAIL_USE_SSL', 'True') == 'True'  
app.config['MAIL_USERNAME'] = os.getenv('MAIL_USERNAME')
app.config['MAIL_PASSWORD'] = os.getenv('MAIL_PASSWORD')
app.config['MAIL_DEFAULT_SENDER'] = os.getenv('MAIL_USERNAME')
app.config['MAIL_TIMEOUT'] = 30  

mail = Mail(app)


limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=["1000 per day", "200 per hour"],
    storage_uri="memory://"
)


import os
basedir = os.path.abspath(os.path.dirname(__file__)) 

database_url = os.getenv("DATABASE_URL")
if database_url and database_url.startswith("postgres://"):
    database_url = database_url.replace("postgres://", "postgresql://", 1)


app.config['SQLALCHEMY_DATABASE_URI'] = database_url or 'sqlite:///' + os.path.join(basedir, 'database.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db.init_app(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "auth"
login_manager.login_message = "Lütfen giriş yapınız."
login_manager.login_message_category = "warning"

with app.app_context():
    try:
        db.create_all()
        
      
        admin = User.query.filter_by(username="admin").first()
        if not admin:
            print("⚠️ Admin kullanıcısı bulunamadı, oluşturuluyor...")
            hashed_pw = generate_password_hash("Admin123!", method='pbkdf2:sha256')
            

            try:
                admin_user = User(
                    username="admin",
                    email="zeyn3pkoc18@gmail.com",
                    password=hashed_pw,
                    is_admin=True 
                )
            except TypeError:

                print("UYARI: 'is_admin' sütunu bulunamadı! Lütfen models.py dosyasını güncelleyin.")
                admin_user = User(
                    username="admin",
                    email="zeyn3pkoc18@gmail.com",
                    password=hashed_pw
                )

            db.session.add(admin_user)
            db.session.commit()
            print("✅✅✅ ADMIN KULLANICISI OLUŞTURULDU! ✅✅✅")
            print("👉 Kullanıcı Adı: admin")
            print("👉 Şifre: Admin123!")
        else:
            print("ℹ️ Admin kullanıcısı zaten mevcut.")

    except Exception as e:
        print(f"❌ Veritabanı Başlatma Hatası: {e}")

@login_manager.user_loader
def load_user(user_id):
    try: return db.session.get(User, int(user_id))
    except: return None

logging.basicConfig(level=logging.INFO, handlers=[logging.FileHandler("app.log"), logging.StreamHandler()])
logger = logging.getLogger(__name__)
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")





def generate_qr_code(data):
    """Verilen metinden QR Kod üretir ve Base64 string döner."""
    qr = qrcode.QRCode(version=1, box_size=5, border=2)
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    
    buffered = io.BytesIO()
    img.save(buffered, format="PNG")
    img_str = base64.b64encode(buffered.getvalue()).decode("utf-8")
    return f"data:image/png;base64,{img_str}"

def clean_json_string(s):
    """
    Yapay zekadan gelen metni temizler.
    ```json ve ``` gibi Markdown etiketlerini siler.
    """
    if s is None:
        return ""
   
    cleaned = s.replace("```json", "").replace("```", "").strip()
    return cleaned


def call_openai_api(system_prompt, user_content="Devam et"):
    if not OPENAI_API_KEY or "benim keyim" in OPENAI_API_KEY:
        return None
    
    payload = {
        "model": "gpt-4o-mini",
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_content}
        ],
        "temperature": 0.1,
        "response_format": {"type": "json_object"}
    }
    
    try:
        r = requests.post("https://api.openai.com/v1/chat/completions", headers={"Authorization": f"Bearer {OPENAI_API_KEY}"}, json=payload, timeout=300)
        if r.status_code == 200:
            return r.json()
        else:
            logger.error(f"OpenAI API Hatası: {r.status_code} - {r.text}")
            return None
    except Exception as e:
        logger.error(f"API Bağlantı Hatası: {e}")
        return None

def get_mock_response(text):
    return {
        "status": "DILEKCE_HAZIR", 
        "dilekce_metni": f"SİMÜLASYON METNİ (API Hatası): {text}", 
        "hukuki_oneriler": "Bağlantı sorunu.", 
        "improvement_note": "Denetim yapılamadı.", 
        "usage": {}
    }



def run_agent_chain(context_data, rag_context):
    start_time = time.time()
    total_tokens = 0
    bugun = datetime.now().strftime("%d.%m.%Y")


    profil_notu = ""
    if current_user.is_authenticated:
        ad_soyad = current_user.full_name or "Belirtilmemiş"
        tckn = current_user.tckn or "Belirtilmemiş"
        adres = current_user.address or "Belirtilmemiş"
        profil_notu = f"""
        [SİSTEM NOTU - KULLANICI KİMLİK KARTI]:
        AD SOYAD: {ad_soyad}
        TCKN: {tckn}
        ADRES: {adres}
        """
        context_data = profil_notu + "\n\n" + context_data 


    question_count = context_data.count("Kullanıcı:")
    

    if question_count >= 5:
        logger.warning("⚠️ Soru limiti aşıldı, Writer'a geçiliyor.")
        
   
        logger.info("🤖 Ajan 1: Avukat taslağı hazırlıyor (Zorla)...")
        prompt1 = WRITER_PROMPT.format(context_data=context_data)
        resp1 = call_openai_api(prompt1)
        
        if not resp1: 
            return get_mock_response("API Hatası")
        

        raw_content1 = resp1["choices"][0]["message"]["content"]
        content1 = json.loads(clean_json_string(raw_content1))
        total_tokens += resp1["usage"]["total_tokens"]
        
        draft_content = content1.get("draft_content", "")
        
  
        logger.info("⚖️ Ajan 2: Savcı/Denetçi inceliyor...")
        prompt2 = CRITIC_PROMPT.format(draft_content=draft_content, rag_context=rag_context)
        resp2 = call_openai_api(prompt2)
        
        critique = "Eleştiri yok."
        if resp2:
            content2 = json.loads(clean_json_string(resp2["choices"][0]["message"]["content"]))
            critique = content2.get("critique", "Mükemmel.")
            total_tokens += resp2["usage"]["total_tokens"]
        

        logger.info("✍️ Ajan 3: Yönetici son hali veriyor...")

        cat_key_match = re.search(r"Kategori:\s*(\w+)", context_data) # [cite: 2738]
        cat_key = cat_key_match.group(1) if cat_key_match else "genel" # [cite: 2742]
        
        prompt3 = FINALIZER_PROMPT.format(draft_content=draft_content, critique=critique, category_key=cat_key)
        resp3 = call_openai_api(prompt3)
        
        final_result = {}
        if resp3:
            try:
                raw_content3 = resp3["choices"][0]["message"]["content"]
                final_result = json.loads(clean_json_string(raw_content3))

                if "dilekce_metni" in final_result:
                    temiz_metin = final_result["dilekce_metni"]
                
                 
                    temiz_metin = re.sub(r'<br\s*/?>', '\n', temiz_metin)
                    temiz_metin = re.sub(r'</?b>', '', temiz_metin)
                    temiz_metin = re.sub(r'</?strong>', '', temiz_metin)
                    temiz_metin = re.sub(r'</?p>', '', temiz_metin)
                    temiz_metin = re.sub(r'</?center>', '', temiz_metin)
               
               
                    temiz_metin = temiz_metin.replace('**', '')
                    temiz_metin = temiz_metin.replace('##', '')
                    temiz_metin = temiz_metin.strip()
               
                    final_result["dilekce_metni"] = temiz_metin

                
                if "graph_data" not in final_result or not final_result["graph_data"]:
                    final_result["graph_data"] = {
                        "nodes": [{"id": 1, "label": "Başvuran", "group": "person"}],
                        "edges": []
                    }
                
                total_tokens += resp3["usage"]["total_tokens"]
            except json.JSONDecodeError:
                final_result = {
                    "status": "DILEKCE_HAZIR", 
                    "dilekce_metni": draft_content, 
                    "hukuki_oneriler": "Format hatası.", 
                    "improvement_note": "Sistem hatası.",
                    "graph_data": {} 
                }
        
        final_result["usage"] = {
            "total_tokens": total_tokens,
            "processing_time": round(time.time() - start_time, 2)
        }
        
        return final_result
    

    logger.info("🕵️ Ajan 0: Mülakatçı eksik bilgi kontrolü yapıyor...")
    

    req_fields = "Tarih, Tutar, Olayın Özeti, Karşı Taraf" 
    cat_match = re.search(r"Kategori:\s*(\w+)", context_data)
    
    kategori_baslik = "Genel Hukuk İşlemleri"
    cat_key = None
    if cat_match:
        cat_key = cat_match.group(1)
        kategori_baslik = CATEGORIES.get(cat_key, {}).get("title", f"Bilinmeyen Kategori ({cat_key})")
        if cat_key in CATEGORIES:
            req_fields = ", ".join(CATEGORIES[cat_key]["required_fields"])




    prompt0 = INTERROGATOR_PROMPT.format(
        context_data=context_data, 
        required_fields=req_fields,
        current_date=bugun,
        kategori=kategori_baslik
    )
    
    resp0 = call_openai_api(prompt0)
    if not resp0: return get_mock_response("API Hatası")
    
    raw_content0 = resp0["choices"][0]["message"]["content"]
    content0 = json.loads(clean_json_string(raw_content0))
    total_tokens += resp0["usage"]["total_tokens"]

    status = content0.get("status", "").upper()
    
    if "TAM" not in status and "YETER" not in status:
        return {
            "status": "SORGU",
            "questions": [str(content0.get("message", "Eksik bilgileri tamamlayalım."))],
            "usage": {"total_tokens": total_tokens}
        }
    

    logger.info("🤖 Ajan 1: Avukat taslağı hazırlıyor...")
    prompt1 = WRITER_PROMPT.format(context_data=context_data)
    resp1 = call_openai_api(prompt1)
    
    if not resp1: return get_mock_response("API Hatası")
    
    raw_content1 = resp1["choices"][0]["message"]["content"]
    content1 = json.loads(clean_json_string(raw_content1))
    total_tokens += resp1["usage"]["total_tokens"]
    
    if content1.get("status") == "SORGU":
        return {
            "status": "SORGU",
            "questions": content1.get("questions", []),
            "usage": {"total_tokens": total_tokens}
        }
    
    draft_content = content1.get("draft_content", "")

    draft_content = re.sub(r'<br\s*/?>', '\n', draft_content)      # <br> → \n
    draft_content = re.sub(r'</?b>', '', draft_content)             # <b> sil
    draft_content = re.sub(r'</?strong>', '', draft_content)        # <strong> sil
    draft_content = re.sub(r'</?p>', '', draft_content)             # <p> sil
    draft_content = re.sub(r'</?center>', '', draft_content)        # <center> sil
    draft_content = draft_content.replace('**', '')                 # ** sil
    draft_content = draft_content.replace('##', '')                 # ## sil
    draft_content = draft_content.strip()

    logger.info(f"📝 Metin temizlendi: {len(draft_content)} karakter")

  
    logger.info("⚖️ Ajan 2: Savcı/Denetçi inceliyor...")
    prompt2 = CRITIC_PROMPT.format(draft_content=draft_content, rag_context=rag_context)
    resp2 = call_openai_api(prompt2)
    
    critique = "Eleştiri yok."
    if resp2:
        content2 = json.loads(clean_json_string(resp2["choices"][0]["message"]["content"]))
        critique = content2.get("critique", "Mükemmel.")
        total_tokens += resp2["usage"]["total_tokens"]
    

    logger.info("✍️ Ajan 3: Yönetici son hali veriyor...")


    cat_key_match = re.search(r"Kategori:\s*(\w+)", context_data)
    cat_key = cat_key_match.group(1) if cat_key_match else "genel"

  
    prompt3 = FINALIZER_PROMPT.format(
        draft_content=draft_content, 
        critique=critique,
        category_key=cat_key 
    )

    resp3 = call_openai_api(prompt3)
    
    final_result = {}
    if resp3:
        try:
            raw_content3 = resp3["choices"][0]["message"]["content"]
            final_result = json.loads(clean_json_string(raw_content3))
            
            if "graph_data" not in final_result or not final_result["graph_data"]:
                final_result["graph_data"] = {
                    "nodes": [{"id": 1, "label": "Başvuran", "group": "person"}],
                    "edges": []
                }
            
            total_tokens += resp3["usage"]["total_tokens"]
        except json.JSONDecodeError:
            final_result = {
                "status": "DILEKCE_HAZIR", 
                "dilekce_metni": draft_content, 
                "hukuki_oneriler": "Format hatası.", 
                "improvement_note": "Sistem hatası.",
                "graph_data": {} 
            }
    else:
        final_result = {
            "status": "DILEKCE_HAZIR", 
            "dilekce_metni": draft_content, 
            "hukuki_oneriler": "Bağlantı hatası.", 
            "improvement_note": "",
            "graph_data": {}
        }

    final_result["usage"] = {
        "total_tokens": total_tokens,
        "processing_time": round(time.time() - start_time, 2)
    }
    
    return final_result

def html_to_pdf_playwright(html_content: str) -> bytes:
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page()
            page.set_content(html_content)
            return page.pdf(format="A4", margin={"top":"0","right":"0","bottom":"0","left":"0"}) 
    except Exception as e:
        logger.error(f"PDF Hatası: {e}")
        raise e

def save_petition_to_db(result, kategori, ocr_data):
    try:
        usage = result.get("usage", {})
        content_text = result.get("dilekce_metni", "")
        hitap_makam_text = result.get("hitap_makam", "DİLEKÇE BAŞLIĞI YOK")

        prediction = success_predictor.analyze_case(kategori, content_text, ocr_data)
        session['prediction'] = prediction
        
    
        doc_hash = hashlib.sha256(content_text.encode('utf-8')).hexdigest()

        new_petition = Petition(
            category=CATEGORIES.get(kategori, {}).get("title", kategori),
            content=content_text,
            hitap_makam=hitap_makam_text,
            advice=result.get("hukuki_oneriler", ""),
            improvement_note=result.get("improvement_note", ""),
            
        
            graph_data=result.get("graph_data", {}), 
         

            ocr_data=ocr_data,
            user_id=current_user.id,
            processing_time=usage.get("processing_time", 0.0),
            token_count=usage.get("total_tokens", 0),
            cost_usd=(usage.get("total_tokens", 0) * 0.00000015),
            status="Taslak",
            tracking_number=f"TR-{random.randint(100000,999999)}",
            document_hash=doc_hash
        )
        db.session.add(new_petition)
        db.session.commit()
        
        session['current_petition_id'] = new_petition.id
        session.pop('result', None)
        session.pop('privacy_map', None)
        session.pop('state', None)
        
        logger.info(f"Dilekçe mühürlendi (Hash: {doc_hash[:8]}...) ve kaydedildi.")
    except Exception as e:
        logger.error(f"DB Kayıt Hatası: {e}")




@app.route("/similar/<int:petition_id>")
@login_required
def similar_petitions(petition_id):
   
    current_petition = db.session.get(Petition, petition_id)
    
    if not current_petition or current_petition.user_id != current_user.id:
        abort(403)


    similar = Petition.query.filter(
        Petition.category == current_petition.category, 
        Petition.id != petition_id,                     
        Petition.user_id == current_user.id             
    ).order_by(Petition.date_created.desc()).limit(5).all()

    return render_template("similar.html", petitions=similar, current=current_petition)

@app.route("/download_word/<int:p_id>")
@login_required
def download_word(p_id):
    """WORD (.docx) ÇIKTISI OLUŞTURMA - KURUMA ÖZEL FORMAT"""
    try:
        if not p_id: 
            return redirect(url_for("dashboard"))
        
        petition = db.session.get(Petition, p_id)
        
     
        if not petition or petition.user_id != current_user.id:
            return redirect(url_for("dashboard")) 
        
        document = Document()
        
  
        match = re.search(r'(<center><b>.*?</b></center><br>)', petition.content, re.IGNORECASE)
        
        hitap_makam = "DİLEKÇE BAŞLIĞI"
        temiz_metin = petition.content
        
        if match:
            makam_html = match.group(1)
            hitap_makam = makam_html.replace('<center><b>', '').replace('</b></center><br>', '').strip()
            temiz_metin = petition.content.replace(makam_html, "").strip()
        
  
        baslik = document.add_heading(hitap_makam, level=1)
        baslik.alignment = 1  
        
        
        p = document.add_paragraph()
        p.add_run(f"Doküman No: {petition.tracking_number}\n").bold = True
        p.add_run(f"Tarih: {datetime.now().strftime('%d.%m.%Y')}\n")
        
  
        doc_hash = petition.document_hash if petition.document_hash else "Mühür Yok"
        p.add_run(f"Dijital Mühür (SHA-256): {doc_hash[:20]}...").italic = True
        
        document.add_paragraph()  
        
 
        temiz_metin = temiz_metin.replace('<br>', '\n').replace('<b>', '').replace('</b>', '')
        temiz_metin = temiz_metin.replace('<strong>', '').replace('</strong>', '')
        
        document.add_paragraph(temiz_metin)
        

        document.add_heading('HUKUKİ DAYANAK VE TAVSİYELER', level=2)
        document.add_paragraph(petition.advice)
        
    
        document.add_paragraph()
        document.add_paragraph(f"Saygılarımla,\n{current_user.full_name or current_user.username}")
        
      
        f = io.BytesIO()
        document.save(f)
        f.seek(0)
        
        return send_file(
            f, 
            as_attachment=True, 
            download_name=f"Dilekce_{p_id}.docx", 
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    except Exception as e:
        logger.error(f"Word Hatası: {e}")
        flash("WORD dosyası oluşturulamadı.", "danger")
        return redirect(url_for("sonuc"))

@app.route("/pdf/<int:p_id>") 
@login_required
@limiter.limit("10 per hour")
def pdf(p_id): 
    try:
        if not p_id: return "Dilekçe ID yok", 404
        
        petition = db.session.get(Petition, p_id)
        if not petition or petition.user_id != current_user.id: 
            return "Yetkisiz Erişim", 403

      
        hitap_makam_h2 = petition.hitap_makam or "DİLEKÇE BAŞLIĞI YOK"
        temiz_dilekce_metni = petition.content
        
        
       
        temiz_dilekce_metni = re.sub(r'<br\s*/?>', '\n', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?b>', '', temiz_dilekce_metni)        # <b> sil
        temiz_dilekce_metni = re.sub(r'</?strong>', '', temiz_dilekce_metni)   # <strong> sil
        temiz_dilekce_metni = re.sub(r'</?p>', '', temiz_dilekce_metni)        # <p> sil
        temiz_dilekce_metni = re.sub(r'</?center>', '', temiz_dilekce_metni)   # <center> sil
        
       
        temiz_dilekce_metni = temiz_dilekce_metni.replace('**', '')
        temiz_dilekce_metni = temiz_dilekce_metni.replace('##', '')
      
    
        qr_data = f"JusticAI|ID: {petition.tracking_number}|HASH: {petition.document_hash}"
        qr_img = generate_qr_code(qr_data)
        qr_html_block = f'<img src="{qr_img}" width="80"><br><span>{petition.tracking_number}</span>'
        
     
        imza_html = f'<img src="{petition.signature_data}" class="signature-img" alt="İmza">' if petition.signature_data else '<p style="color:#ccc;">(İmza)</p>'
        
        
        html = PETITION_HTML_TEMPLATE.format(
            hitap_makam=hitap_makam_h2,
            dilekce_metni=temiz_dilekce_metni.replace("\n", "<br>"),  # Satır sonlarını <br> yap
            hukuki_oneriler=petition.advice.replace("\n", "<br>"),
            improvement_note=petition.improvement_note or "Standart denetim.",
            tracking_number=petition.tracking_number or "TASLAK",
            tarih=datetime.now().strftime("%d.%m.%Y"),
            ad_soyad=current_user.full_name or current_user.username,
            imza_resmi=imza_html,
            qr_kod_html=qr_html_block,
            document_hash=petition.document_hash or "YOK"
        )
        
      
        pdf_bytes = html_to_pdf_playwright(html)
        return send_file(io.BytesIO(pdf_bytes), as_attachment=True, download_name=f"Dilekce_{p_id}.pdf", mimetype="application/pdf")
        
    except Exception as e:
        logger.error(f"PDF Hatası: {e}")
        return f"PDF Hatası: {str(e)}", 500

@app.route("/", methods=["GET", "POST"])
@login_required
@limiter.limit("10 per minute")
def index():
    try:
        if request.method == "GET":
            session.pop('privacy_map', None)
            session.pop('result', None)
            return render_template("index.html", categories=CATEGORIES, mode="initial", user=current_user)
        
        step = request.form.get("step")
        if step == "answers": return cevap_ver()
        
        ocr_text = ""
        uploaded_file = request.files.get("dosya")
        if uploaded_file and uploaded_file.filename != '':
            res = extract_text_from_file(uploaded_file)
            
           
            if "error" in res:
                flash(f"❌ OCR Hatası: {res['error']}", "danger")
              
                if "suggestion" in res:
                    flash(f"💡 İpucu: {res['suggestion']}", "info")
                return redirect(url_for("index"))
          

            if "text" in res: 
                ocr_text = res["text"]

        kategori = request.form.get("kategori")
        aciklama = request.form.get("aciklama", "").strip()
        
        if not aciklama and not ocr_text:
            flash("Lütfen bir açıklama yazın.", "danger")
            return redirect(url_for("index"))

        is_ethical, ethics_msg = ethics_guard.analyze(aciklama + " " + ocr_text)
        if not is_ethical:
            flash(f"GÜVENLİK UYARISI: {ethics_msg}", "danger")
            return redirect(url_for("index"))

        rules_feedback = check_rules(kategori, aciklama, ocr_text)
        rag_context = search_legal_docs(aciklama + " " + ocr_text, kategori) or "Genel hukuk kuralları ve Türk Borçlar Kanunu hükümleri geçerlidir."
        
        raw_combined = f"KULLANICI AÇIKLAMASI:\n{aciklama}\n\n=== YÜKLENEN BELGE İÇERİĞİ (OCR) ===\n{ocr_text}\n======================================"
        
        safe_text, privacy_mapping = privacy_guard.anonymize(raw_combined)
        session['privacy_map'] = privacy_mapping

     
        result = run_agent_chain(
            context_data=f"Kategori: {kategori}\nKullanıcı Beyanı: {safe_text}\nRiskler: {rules_feedback}", 
            rag_context=rag_context
        )
        
        if result.get("status") == "DILEKCE_HAZIR":
            restored = privacy_guard.de_anonymize(result.get("dilekce_metni", ""), privacy_mapping)
            result["dilekce_metni"] = restored
            save_petition_to_db(result, kategori, ocr_text)

        if result.get("status") == "SORGU":
            session["state"] = {"kategori": kategori, "aciklama": safe_text}
            return render_template("index.html", mode="questions", sorular=list(enumerate(result.get("questions", []))), ad_soyad=current_user.username)

        return redirect(url_for("sonuc"))

    except Exception as e:
        logger.error(f"Index Error: {e}")
        flash("Hata oluştu.", "danger")
        return redirect(url_for("index"))

def cevap_ver():
    try:
        state = session.get("state")
        if not state: 
           
            return redirect(url_for("index"))
        
        question_count = state.get("question_count", 0)
        state["question_count"] = question_count + 1  
        
        privacy_mapping = session.get('privacy_map', {})
        
        cevaplar = "\n".join([f"- {v}" for k,v in request.form.items() if k.startswith("cevap_")])

        gecmis_hikaye = state.get("aciklama", "")
        

        if "Kategori:" not in gecmis_hikaye:
             gecmis_hikaye = f"Kategori: {state.get('kategori', 'Genel')}\n" + gecmis_hikaye

        guncel_context = f"{gecmis_hikaye}\n\n[KULLANICI CEVABI]:\n{cevaplar}"
        
        
        result = run_agent_chain(
            context_data=guncel_context,
            rag_context="Mevcut sohbet devam ediyor." 
        )
        
        if result.get("status") == "SORGU":
             questions = result.get("questions", [])
  
             if not questions:
                 print("⚠️ AI soru modunda ama soru üretmedi. Dilekçeye zorlanıyor.")

                 pass 
             else:
               
                 state["aciklama"] = guncel_context
                 session["state"] = state
                 session.modified = True
                 

                 return render_template("index.html", 
                                      mode="questions", 
                                      sorular=list(enumerate(questions)), 
                                      ad_soyad=current_user.username)
        
       
        if "dilekce_metni" in result:

            restored_text = privacy_guard.de_anonymize(result["dilekce_metni"], privacy_mapping)
            result["dilekce_metni"] = restored_text
            
  
            save_petition_to_db(result, state["kategori"], "Soru-Cevap Modu")
            

            return redirect(url_for("sonuc"))
            
      
        return redirect(url_for("index"))
        
    except Exception as e:
        print(f"Cevap Ver Hatası: {e}")
        return redirect(url_for("index"))

@app.route("/save_signature", methods=["POST"])
@login_required
def save_signature():
    try:
        data = request.json
        p_id = session.get('current_petition_id')
        if not p_id: return {"status": "error"}, 404
        petition = db.session.get(Petition, p_id)
        if petition and petition.user_id == current_user.id:
            petition.signature_data = data.get("signature")
            db.session.commit()
            return {"status": "success"}
        return {"status": "error"}, 403
    except: return {"status": "error"}, 500

@app.route("/download_xml/<int:p_id>") 
@login_required
def download_xml(p_id): 
    try:
       
        if not p_id: return redirect(url_for("dashboard"))
        
        petition = db.session.get(Petition, p_id)
        
        if not petition or petition.user_id != current_user.id:
            return redirect(url_for("dashboard"))

        
        xml_content = f"""<?xml version="1.0"?><UYAP><Icerik>{petition.content}</Icerik><Hash>{petition.document_hash}</Hash></UYAP>"""
        
        mem = io.BytesIO()
        mem.write(xml_content.encode('utf-8'))
        mem.seek(0)
        
        return send_file(mem, mimetype='application/xml', as_attachment=True, download_name=f"UYAP_{p_id}.xml")
    except: return redirect(url_for("sonuc"))

@app.route("/sonuc")
@login_required
def sonuc():
    try:

        p_id = request.args.get('id') or session.get('current_petition_id')
        if not p_id: 
            return redirect(url_for("index"))

        petition = db.session.get(Petition, p_id)
        

        if petition and petition.user_id == current_user.id:
            session['current_petition_id'] = petition.id
            
            hitap_makam_h2 = petition.hitap_makam or "DİLEKÇE BAŞLIĞI YOK"
            temiz_dilekce_metni = petition.content
            
            
            temiz_dilekce_metni = re.sub(r'<br\s*/?>', '\n', temiz_dilekce_metni)
            temiz_dilekce_metni = re.sub(r'</?b>', '', temiz_dilekce_metni)
            temiz_dilekce_metni = re.sub(r'</?strong>', '', temiz_dilekce_metni)
            temiz_dilekce_metni = re.sub(r'</?p>', '', temiz_dilekce_metni)
            temiz_dilekce_metni = re.sub(r'</?center>', '', temiz_dilekce_metni)
            temiz_dilekce_metni = temiz_dilekce_metni.replace('**', '').replace('##', '')            
            

            graph_payload = petition.graph_data or {}
            
            data = {
                "hitap_makam": hitap_makam_h2,
                "dilekce_metni": temiz_dilekce_metni, 
                "hukuki_oneriler": petition.advice,
                "graph_data": graph_payload
            }
            return render_template("sonuc.html", data=data, ad_soyad=current_user.username)
        
        return redirect(url_for("index"))
    except Exception as e: 
        print(f"Sonuç Hatası: {e}")
        return redirect(url_for("index"))

@app.route("/send_mail")
@login_required
@limiter.limit("3 per hour")
def send_mail():
    try:
        if not app.config['MAIL_USERNAME']: return redirect(url_for("sonuc"))
        p_id = session.get('current_petition_id')
        petition = db.session.get(Petition, p_id)

        if not petition or petition.user_id != current_user.id: 
            flash("Yetkisiz Erişim", "danger")
            return redirect(url_for("sonuc"))
        
        hitap_makam_h2 = petition.hitap_makam or "DİLEKÇE BAŞLIĞI YOK"
        temiz_dilekce_metni = petition.content
        
        temiz_dilekce_metni = re.sub(r'<br\s*/?>', '\n', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?b>', '', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?strong>', '', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?p>', '', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?center>', '', temiz_dilekce_metni)
        temiz_dilekce_metni = temiz_dilekce_metni.replace('**', '').replace('##', '')

        msg = Message(subject=f"JusticAI - {hitap_makam_h2} Başvurusu", recipients=[current_user.email], body="Ektedir.")
        
        imza_html = f'<img src="{petition.signature_data}" class="imza-resmi" alt="İmza">' if petition.signature_data else '<p>(İmza)</p>'

        if petition.tracking_number and petition.document_hash:
            qr_html_block = f'''
            <img src="https://api.qrserver.com/v1/create-qr-code/?size=80x80&data={petition.tracking_number}" alt="QR">
            <span>Takip No: {petition.tracking_number[:8]}</span>
            '''
        else:
            qr_html_block = '<span style="font-size:10px; color:#999;">Taslak</span>'
       
        html = PETITION_HTML_TEMPLATE.format(
            hitap_makam=hitap_makam_h2, 
            dilekce_metni=temiz_dilekce_metni.replace("\n", "<br>"),
            hukuki_oneriler=petition.advice.replace("\n", "<br>"),
            improvement_note=petition.improvement_note or "",
            tracking_number=petition.tracking_number or "",
            tarih=datetime.now().strftime("%d.%m.%Y"),
            ad_soyad=current_user.username,
            imza_resmi=imza_html,
            qr_kod_html=qr_html_block,
            document_hash=petition.document_hash or "YOK"  
        )
        pdf_bytes = html_to_pdf_playwright(html)
        
        msg.attach("Dilekce.pdf", "application/pdf", pdf_bytes)
        mail.send(msg)
        flash("✅ Dilekçeniz e-posta ile başarıyla gönderildi!", "success")
        return redirect(url_for("sonuc"))
    except Exception as e: 
        logger.error(f"Mail Gönderme Hatası: {e}")
        flash("❌ E-posta gönderme sırasında bir hata oluştu.", "danger")
        return redirect(url_for("sonuc"))

@app.route("/submit_feedback", methods=["POST"])
@login_required
def submit_feedback():
    try:
        data = request.json
        p_id = session.get('current_petition_id')
        
       
        if not p_id: 
            return {"status": "error", "message": "Aktif dilekçe bulunamadı."}, 404


        existing = Feedback.query.filter_by(petition_id=p_id, user_id=current_user.id).first()
        
        if existing:
            return {"status": "error", "message": "Bu dilekçe için zaten geri bildirim verdiniz."}, 400

     
        new_f = Feedback(
            petition_id=p_id,
            user_id=current_user.id, 
            rating=data.get("rating"),
            comment=data.get("comment")
        )
        db.session.add(new_f)
        db.session.commit()
        
        logger.info(f"Feedback alındı: {'Olumlu' if data.get('rating') else 'Olumsuz'} - User: {current_user.username}")
        return {"status": "success", "message": "Geri bildiriminiz için teşekkürler!"}

    except Exception as e:
        logger.error(f"Feedback Hatası: {e}")
        return {"status": "error", "message": str(e)}, 500

@app.route("/auth", methods=["GET"])
def auth(): return render_template("auth.html", active_tab="login")

@app.route("/register", methods=["POST"])
def register():
    try:
        username = request.form.get("username")
        password = request.form.get("password")
        email = request.form.get("email")
        if not email or User.query.filter_by(username=username).first(): return render_template("auth.html", active_tab="register")
        new_user = User(username=username, email=email, password=generate_password_hash(password, method='pbkdf2:sha256'))
        db.session.add(new_user); db.session.commit()
        return render_template("auth.html", active_tab="login")
    except: return render_template("auth.html", active_tab="register")

@app.route("/login", methods=["GET", "POST"])
def login():
    try:
        if request.method == "GET": return render_template("auth.html", active_tab="login")
        user = User.query.filter_by(username=request.form.get("username")).first()
        if user and check_password_hash(user.password, request.form.get("password")):
            login_user(user); return redirect(url_for("index"))
        return render_template("auth.html", active_tab="login")
    except: return render_template("auth.html", active_tab="login")

@app.route("/logout")
def logout():
    logout_user()
    return redirect(url_for("auth"))

@app.route("/dashboard")
@login_required
def dashboard():
    user_petitions = Petition.query.filter_by(user_id=current_user.id).order_by(Petition.date_created.desc()).all()
    return render_template("dashboard.html", petitions=user_petitions)

@app.route("/admin")
@login_required
def admin_panel():
    if not current_user.is_admin:
        flash("⛔ Bu sayfaya erişim yetkiniz yok.", "danger")
        return redirect(url_for("dashboard"))
    
    feedbacks = Feedback.query.order_by(Feedback.created_at.desc()).limit(20).all()
    total_feedbacks = Feedback.query.count()
    positive = Feedback.query.filter_by(rating=True).count()
    satisfaction_rate = int((positive / total_feedbacks) * 100) if total_feedbacks > 0 else 0
    
    total_users = User.query.count()
    total_petitions = Petition.query.count()
    category_stats = db.session.query(Petition.category, func.count(Petition.id)).group_by(Petition.category).all()
    avg_time = db.session.query(func.avg(Petition.processing_time)).scalar() or 0.0
    total_tokens = db.session.query(func.sum(Petition.token_count)).scalar() or 0
    total_cost = db.session.query(func.sum(Petition.cost_usd)).scalar() or 0.0
    
    return render_template("admin.html", total_users=total_users, total_petitions=total_petitions, 
                           labels=[c[0] for c in category_stats], data=[c[1] for c in category_stats], 
                           avg_time=round(avg_time, 2), total_tokens=total_tokens, total_cost=round(total_cost, 5),
                           feedbacks=feedbacks, satisfaction_rate=satisfaction_rate, total_feedbacks=total_feedbacks)

@app.route("/edit/<int:petition_id>", methods=["GET", "POST"])
@login_required
def edit_petition(petition_id):
    """Dilekçe düzenleme route'u - dashboard.html'de kullanılıyor"""
 
    petition = db.session.get(Petition, petition_id)
    

    if not petition or petition.user_id != current_user.id:
        abort(403)
    
   
    if petition.status == "İletildi":
        flash("⚠️ Gönderilmiş dilekçeler düzenlenemez!", "warning")
        return redirect(url_for("dashboard"))
    
    if request.method == "POST":
        new_content = request.form.get("content")
        
     
        new_hash = hashlib.sha256(new_content.encode('utf-8')).hexdigest()
        
        petition.content = new_content
        petition.document_hash = new_hash
        
        db.session.commit()
        flash("✅ Dilekçe güncellendi ve yeni dijital mühür oluşturuldu.", "success")
        return redirect(url_for("dashboard"))
    
    return render_template("edit.html", petition=petition)

@app.route("/simulate_edevlet/<int:petition_id>")
@login_required
def simulate_edevlet(petition_id):
    """e-Devlet simülasyonu - dashboard.html'de kullanılıyor"""
    petition = db.session.get(Petition, petition_id)
    
    if not petition or petition.user_id != current_user.id:
        abort(403)
    

    petition.status = "İletildi"
    petition.tracking_number = f"EDEVLET-{random.randint(100000, 999999)}"
    db.session.commit()
    
    flash("✅ Dilekçe e-Devlet sistemine (simüle) başarıyla iletildi!", "success")
    return redirect(url_for("dashboard"))

@app.route("/run_tests_web")
@login_required
def run_tests_web():
    if not current_user.is_admin:
        flash("⛔ Bu sayfaya erişim yetkiniz yok.", "danger")
        return redirect(url_for("dashboard"))
    log_capture_string = io.StringIO()
    runner = unittest.TextTestRunner(stream=log_capture_string, verbosity=2)
    result = runner.run(unittest.TestLoader().loadTestsFromModule(tests))
    output = log_capture_string.getvalue()
    final_status = "TÜM SİSTEMLER OPERASYONEL 🚀" if result.wasSuccessful() else "SİSTEMDE SORUNLAR TESPİT EDİLDİ ⚠️"
    color = "text-success" if result.wasSuccessful() else "text-danger"
    return f"<div class='p-3 mb-3 border rounded bg-dark text-light font-monospace' style='font-size: 0.85rem; white-space: pre-wrap;'>{output}</div><h5 class='fw-bold {color} mt-3'>{final_status}</h5>"

@app.route("/get_pdf_content")
@login_required
def get_pdf_content():
    try:
        p_id = session.get('current_petition_id')
        if not p_id: return abort(404)
        petition = db.session.get(Petition, p_id)

        if not petition or petition.user_id != current_user.id: 
            return abort(403)
        
        hitap_makam_h2 = petition.hitap_makam or "DİLEKÇE BAŞLIĞI YOK"
        temiz_dilekce_metni = petition.content

        temiz_dilekce_metni = re.sub(r'<br\s*/?>', '\n', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?b>', '', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?strong>', '', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?p>', '', temiz_dilekce_metni)
        temiz_dilekce_metni = re.sub(r'</?center>', '', temiz_dilekce_metni)
        temiz_dilekce_metni = temiz_dilekce_metni.replace('**', '').replace('##', '')
     

        html = PETITION_HTML_TEMPLATE.format(
            hitap_makam=hitap_makam_h2, 
            dilekce_metni=temiz_dilekce_metni.replace("\n", "<br>"), # <--- TEMİZ GÖVDE
            hukuki_oneriler="", 
            tarih=datetime.now().strftime("%d.%m.%Y"),
            ad_soyad=current_user.username,
            imza_resmi="",
            improvement_note="", tracking_number="" 
        )
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page()
            page.set_content(html)
            pdf_bytes = page.pdf(format="A4", margin={"top":"0","right":"0","bottom":"0","left":"0"})
        return send_file(io.BytesIO(pdf_bytes), mimetype='application/pdf', as_attachment=False, download_name='preview.pdf')
    except Exception as e:
        logger.error(f"PDF Önizleme Hatası: {e}")
        return abort(500)

@app.template_filter('strftime')
def _jinja2_filter_strftime(date, fmt=None):
    if not date or date == "now":
        return datetime.now().strftime(fmt or '%d.%m.%Y')
    if hasattr(date, 'strftime'):
        return date.strftime(fmt or '%d.%m.%Y')
    return str(date)

@app.errorhandler(429)
def ratelimit_handler(e):
    return "<h1>Çok Fazla İstek Yaptınız!</h1><p>Lütfen biraz bekleyin.</p>", 429



@app.route("/chat_api", methods=["POST"])
@login_required
def chat_api():
    try:
        data = request.json
        user_input = data.get('message')
        
        state = session.get("state")
        if not state:
            return {"status": "error", "message": "Oturum düştü"}

        question_count = state.get("question_count", 0)
        
    
        if question_count >= 5:
            result = run_agent_chain(
                context_data=f"{state.get('aciklama', '')}\nKullanıcı: {user_input}",
                rag_context="Sohbet sınırı aşıldı."
            )
            
           
            if "dilekce_metni" not in result:
                return {"status": "error", "message": "Dilekçe oluşturulamadı. Lütfen tekrar deneyin."}
            
            privacy_mapping = session.get('privacy_map', {})
            restored_text = privacy_guard.de_anonymize(result["dilekce_metni"], privacy_mapping)
            result["dilekce_metni"] = restored_text
            save_petition_to_db(result, state["kategori"], "Chat Modu")
            
            return {
                "status": "finished",
                "redirect": url_for("sonuc"),
                "reply": "Dilekçeniz hazır!"
            }


        gecmis_hikaye = state.get("aciklama", "")
        if "Kategori:" not in gecmis_hikaye:
            gecmis_hikaye = f"Kategori: {state.get('kategori', 'Genel')}\n{gecmis_hikaye}"

        guncel_context = f"{gecmis_hikaye}\nKullanıcı: {user_input}"
        
        result = run_agent_chain(
            context_data=guncel_context,
            rag_context="Sohbet devam ediyor."
        )

        status = result.get("status")
        
        if status == "SORGU":
            questions = result.get("questions", [])
            new_question = questions[0] if questions else "Devam edelim..."
            
            guncel_context += f"\nAsistan: {new_question}"
            state["aciklama"] = guncel_context
            state["question_count"] = question_count + 1
            session["state"] = state
            session.modified = True
            
            return {
                "status": "continue", 
                "reply": new_question
            }
        
      
        if "dilekce_metni" in result:
            privacy_mapping = session.get('privacy_map', {})
            restored_text = privacy_guard.de_anonymize(result["dilekce_metni"], privacy_mapping)
            result["dilekce_metni"] = restored_text
            save_petition_to_db(result, state["kategori"], "Chat Modu")
            session.modified = True
            
            return {
                "status": "finished",
                "redirect": url_for("sonuc"),
                "reply": "Dilekçeniz hazır!"
            }
        
        
        return {"status": "error", "message": "Beklenmeyen bir durum oluştu."}

    except Exception as e:
        print(f"Chat API Hatası: {e}")
        return {"status": "error", "message": str(e)}


@app.route("/predict_category", methods=["POST"])
@login_required
def predict_category():
    """
    GELİŞMİŞ KATEGORİ TAHMİN MOTORU (v2.0)
    Kullanıcı metnini analiz eder ve Mahkeme/Okul/Banka/Belediye ayrımını yapar.
    """
    try:
        data = request.json
        text = data.get("text", "")
    
        if len(text) < 5: return {"category": None}

     
        cat_keys = list(CATEGORIES.keys())
       
        cat_info = "\n".join([f"- {k}: {v['title']}" for k, v in CATEGORIES.items()])

        system_prompt = f"""
        Sen uzman bir Hukuk ve İşlem Sınıflandırma Motorusun. 
        Görevin: Vatandaşın yazdığı sorunu analiz edip, aşağıdaki listeden EN DOĞRU işlem kodunu seçmek.
        
        KATEGORİ LİSTESİ:
        {cat_info}

        KARAR KURALLARI (İPUÇLARI):
        1. **EĞİTİM SİNYALLERİ:** "Hocam", "Not", "Sınav", "Vize", "Büt", "Kayıt", "Ders", "Okul", "Yurt", "KYK" geçerse -> EĞİTİM kategorilerine odaklan.
        2. **FİNANS SİNYALLERİ:** "Kart", "Banka", "Çekim", "Harcama", "İade", "Bloke" geçerse -> BANKA kategorilerine odaklan.
        3. **BELEDİYE SİNYALLERİ:** "Çukur", "Yol", "Çöp", "Lamba", "İmar", "Ruhsat", "Zabıta" geçerse -> BELEDİYE kategorilerine odaklan.
        4. **ABONELİK SİNYALLERİ:** "İnternet", "Kapatmak", "Taahhüt", "Fatura", "Sayaç" geçerse -> ABONELİK kategorilerine odaklan.
        5. **YARGI SİNYALLERİ:** "Boşanma", "Dava", "Şikayet", "Savcılık", "İcra" geçerse -> İlgili YARGI kategorisini seç.

        ÇIKTI FORMATI:
        Sadece JSON döndür: {{"category": "secilen_kategori_kodu"}}
        Eğer emin değilsen veya uygun kategori yoksa: {{"category": null}}
        """

        resp = call_openai_api(system_prompt, user_content=f"Kullanıcı Sorunu: {text}")
        
        if resp:
            raw_content = resp["choices"][0]["message"]["content"]
            result = json.loads(clean_json_string(raw_content))
            return {"category": result.get("category")}
            
        return {"category": None}

    except Exception as e:
        print(f"Kategori Tahmin Hatası: {e}")
        return {"category": None}

@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
   
    if request.method == "POST":
        current_user.full_name = request.form.get("full_name")
        current_user.tckn = request.form.get("tckn")
        current_user.address = request.form.get("address")
        
        db.session.commit() 
        flash("Profil bilgileriniz başarıyla güncellendi!", "success")
        return redirect(url_for("profile"))
    
 
    return render_template("profile.html", user=current_user)

if __name__ == "__main__":
    app.run(debug=True, port=5000)